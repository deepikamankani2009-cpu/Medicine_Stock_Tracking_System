from datetime import datetime
import os
import tkinter as tk
from tkinter import messagebox, ttk
import openpyxl

EXCEL_FILE = "medicine_stock.xlsx"


# --- Excel Initialization ---
def init_excel():
  if not os.path.exists(EXCEL_FILE):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "MedicineStock"
    headers = [
        "Medicine ID",
        "Medicine Name",
        "Category",
        "Quantity",
        "Price (Rs)",
        "Expiry Date",
        "Status",
    ]
    ws.append(headers)
    wb.save(EXCEL_FILE)


# --- Helper Function: Check Status based on Expiry & Stock ---
def get_status(quantity, expiry_str):
  try:
    qty = int(quantity)
    if qty <= 0:
      return "Out of Stock"

    exp_date = datetime.strptime(expiry_str, "%d-%m-%Y")
    today = datetime.now()

    if exp_date < today:
      return "Expired"
    elif (exp_date - today).days <= 30:
      return "Expiring Soon"
    else:
      return "Available"
  except ValueError:
    return "Unknown"


# --- Main Application Class ---
class MedicineApp:

  def __init__(self, root):
    self.root = root
    self.root.title("Medicine Stock & Expiry Tracking System")
    self.root.geometry("900x600")
    self.root.configure(bg="#f4f6f9")

    init_excel()

    # Title Header
    title_label = tk.Label(
        self.root,
        text="💊 Medicine Stock & Expiry Tracking System",
        font=("Arial", 18, "bold"),
        bg="#2c3e50",
        fg="white",
    )
    title_label.pack(fill=tk.X, ipady=10)

    # Form Frame
    form_frame = tk.LabelFrame(
        self.root,
        text=" Add / Manage Medicine ",
        font=("Arial", 12, "bold"),
        bg="#f4f6f9",
        padx=15,
        pady=15,
    )
    form_frame.pack(fill=tk.X, padx=20, pady=10)

    # Form Inputs
    tk.Label(
        form_frame,
        text="Medicine Name:",
        font=("Arial", 10, "bold"),
        bg="#f4f6f9",
    ).grid(row=0, column=0, sticky="w", pady=5)
    self.entry_name = tk.Entry(form_frame, font=("Arial", 10), width=20)
    self.entry_name.grid(row=0, column=1, padx=10, pady=5)

    tk.Label(
        form_frame, text="Category:", font=("Arial", 10, "bold"), bg="#f4f6f9"
    ).grid(row=0, column=2, sticky="w", pady=5)
    self.combo_cat = ttk.Combobox(
        form_frame,
        values=["Tablet", "Capsule", "Syrup", "Injection", "Ointment", "Other"],
        state="readonly",
        width=18,
    )
    self.combo_cat.grid(row=0, column=3, padx=10, pady=5)
    self.combo_cat.current(0)

    tk.Label(
        form_frame, text="Quantity:", font=("Arial", 10, "bold"), bg="#f4f6f9"
    ).grid(row=1, column=0, sticky="w", pady=5)
    self.entry_qty = tk.Entry(form_frame, font=("Arial", 10), width=20)
    self.entry_qty.grid(row=1, column=1, padx=10, pady=5)

    tk.Label(
        form_frame, text="Price (Rs):", font=("Arial", 10, "bold"), bg="#f4f6f9"
    ).grid(row=1, column=2, sticky="w", pady=5)
    self.entry_price = tk.Entry(form_frame, font=("Arial", 10), width=20)
    self.entry_price.grid(row=1, column=3, padx=10, pady=5)

    tk.Label(
        form_frame,
        text="Expiry Date (DD-MM-YYYY):",
        font=("Arial", 10, "bold"),
        bg="#f4f6f9",
    ).grid(row=2, column=0, sticky="w", pady=5)
    self.entry_expiry = tk.Entry(form_frame, font=("Arial", 10), width=20)
    self.entry_expiry.grid(row=2, column=1, padx=10, pady=5)

    # Buttons
    btn_frame = tk.Frame(form_frame, bg="#f4f6f9")
    btn_frame.grid(row=2, column=2, columnspan=2, pady=10)

    tk.Button(
        btn_frame,
        text="Add Medicine",
        command=self.add_medicine,
        bg="#27ae60",
        fg="white",
        font=("Arial", 10, "bold"),
    ).pack(side=tk.LEFT, padx=5)
    tk.Button(
        btn_frame,
        text="Clear Form",
        command=self.clear_entries,
        bg="#7f8c8d",
        fg="white",
        font=("Arial", 10, "bold"),
    ).pack(side=tk.LEFT, padx=5)

    # Table Display (Treeview)
    tree_frame = tk.Frame(self.root, bg="#f4f6f9")
    tree_frame.pack(fill=tk.BOTH, expand=True, padx=20, pady=10)

    cols = (
        "ID",
        "Medicine Name",
        "Category",
        "Quantity",
        "Price (Rs)",
        "Expiry Date",
        "Status",
    )
    self.tree = ttk.Treeview(tree_frame, columns=cols, show="headings", height=10)

    for col in cols:
      self.tree.heading(col, text=col)
      self.tree.column(col, width=120, anchor="center")

    self.tree.pack(fill=tk.BOTH, expand=True)

    # Refresh Data Table
    self.load_data()

  # --- Clear Inputs ---
  def clear_entries(self):
    self.entry_name.delete(0, tk.END)
    self.entry_qty.delete(0, tk.END)
    self.entry_price.delete(0, tk.END)
    self.entry_expiry.delete(0, tk.END)

  # --- Add Medicine to Excel ---
  def add_medicine(self):
    name = self.entry_name.get().strip()
    cat = self.combo_cat.get().strip()
    qty = self.entry_qty.get().strip()
    price = self.entry_price.get().strip()
    exp = self.entry_expiry.get().strip()

    if not name or not qty or not price or not exp:
      messagebox.showwarning("Validation Error", "All fields are required!")
      return

    try:
      datetime.strptime(exp, "%d-%m-%Y")
    except ValueError:
      messagebox.showerror(
          "Invalid Date",
          "Expiry Date format must be DD-MM-YYYY (e.g., 25-12-2026)",
      )
      return

    status = get_status(qty, exp)

    wb = openpyxl.load_workbook(EXCEL_FILE)
    ws = wb["MedicineStock"]

    med_id = f"MED{ws.max_row}"  # ID Generation
    ws.append([med_id, name, cat, int(qty), float(price), exp, status])
    wb.save(EXCEL_FILE)

    messagebox.showinfo("Success", f"Medicine '{name}' added successfully!")
    self.clear_entries()
    self.load_data()

  # --- Read & Load Data from Excel ---
  def load_data(self):
    for item in self.tree.get_children():
      self.tree.delete(item)

    if os.path.exists(EXCEL_FILE):
      wb = openpyxl.load_workbook(EXCEL_FILE)
      ws = wb["MedicineStock"]

      for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0]:  # Ensure ID exists
          updated_status = get_status(row[3], str(row[5]))
          self.tree.insert(
              "",
              tk.END,
              values=(
                  row[0],
                  row[1],
                  row[2],
                  row[3],
                  row[4],
                  row[5],
                  updated_status,
              ),
          )


if __name__ == "__main__":
  root = tk.Tk()
  app = MedicineApp(root)
  root.mainloop()